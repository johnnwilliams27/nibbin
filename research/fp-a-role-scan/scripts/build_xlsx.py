#!/usr/bin/env python3
import csv
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter

FIT_FILL={'5':'C6EFCE','4':'D9EAD3','3':'FFF2CC','2':'FCE5CD','1':'F4CCCC'}
HEADER_FILL=PatternFill('solid',fgColor='1F3864')
HEADER_FONT=Font(bold=True,color='FFFFFF',size=11)
thin=Side(style='thin',color='D9D9D9'); border=Border(left=thin,right=thin,top=thin,bottom=thin)

def add_sheet(wb,title,csv_path,note=None):
    ws=wb.create_sheet(title)
    rows=list(csv.reader(open(csv_path)))
    if not rows: return
    hdr=rows[0]; data=rows[1:]
    link_col=hdr.index('ats_link') if 'ats_link' in hdr else -1
    fit_col=hdr.index('fit_score') if 'fit_score' in hdr else -1
    r0=1
    if note:
        ws.cell(row=1,column=1,value=note).font=Font(italic=True,color='808080',size=10)
        r0=2
    # header
    for j,h in enumerate(hdr,1):
        c=ws.cell(row=r0,column=j,value=h.replace('_',' ')); c.fill=HEADER_FILL; c.font=HEADER_FONT
        c.alignment=Alignment(horizontal='center',vertical='center',wrap_text=True); c.border=border
    # data
    for i,row in enumerate(data,1):
        rr=r0+i
        for j,val in enumerate(row,1):
            c=ws.cell(row=rr,column=j); c.border=border; c.alignment=Alignment(vertical='top',wrap_text=(j in (hdr.index('role_title')+1,hdr.index('domain_note')+1)))
            if j-1==link_col and val:
                c.value='apply ↗'; c.hyperlink=val; c.font=Font(color='0563C1',underline='single')
            else:
                c.value=val
        if fit_col>=0:
            fv=row[fit_col]
            fill=FIT_FILL.get(fv)
            if fill:
                ws.cell(row=rr,column=fit_col+1).fill=PatternFill('solid',fgColor=fill)
                ws.cell(row=rr,column=fit_col+1).font=Font(bold=True)
                ws.cell(row=rr,column=fit_col+1).alignment=Alignment(horizontal='center')
    # widths
    widths={'fit_score':5,'fund(s)':22,'company':20,'ownership':13,'role_title':40,'seniority':14,
            'years_req':8,'remote_scope':22,'posted_date':11,'comp_range':20,'domain_note':34,'ats_link':9}
    for j,h in enumerate(hdr,1):
        ws.column_dimensions[get_column_letter(j)].width=widths.get(h,16)
    ws.freeze_panes=ws.cell(row=r0+1,column=1)
    ws.auto_filter.ref=f"A{r0}:{get_column_letter(len(hdr))}{r0+len(data)}"
    return ws

wb=Workbook(); wb.remove(wb.active)
add_sheet(wb,'Matches','fp_a_remote_roles.csv',
          'Remote-US or Dallas–Fort Worth FP&A / Strategic Finance / Financial-Analyst roles • 8 yrs of experience qualifies (stated min ≤8) • posted ≤60 days • sorted by fit (5=payer core … 1=no overlap) then recency')
add_sheet(wb,'Verify (remote unconfirmed)','fp_a_verify.csv',
          'ATS flags these remote-eligible but the posting lists an office HQ and the JD does not confirm remote — verify the work model before applying.')
wb.save('fp_a_remote_roles.xlsx')
print("wrote fp_a_remote_roles.xlsx with sheets:", wb.sheetnames)
